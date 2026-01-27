# app/views/extruder_benchmark_report/exporter.py

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment


class ExtruderBenchmarkExporter:
    def export(self, summary_valid: pd.DataFrame, summary_invalid: pd.DataFrame, raw_df: pd.DataFrame, filepath: str):
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:

            # --- SHEET 1: RAW DATA ---
            raw_cols = [
                "reference_number", "machine_number", "product_code", "lot_number", "formula_number",
                "time_start", "time_end", "duration_hours", "output_per_hour", "total_output",
                "expected_output", "yield_value",
                "purging_start_time", "purging_end_time", "purging_duration_minutes",
                "total_cleaning_material", "operator", "purging_to_code"
            ]
            valid_raw_cols = [c for c in raw_cols if c in raw_df.columns]
            raw_export = raw_df[valid_raw_cols].copy()

            # Fix Timezones
            date_cols = ["time_start", "time_end", "purging_start_time", "purging_end_time"]
            for col in date_cols:
                if col in raw_export.columns:
                    raw_export[col] = pd.to_datetime(raw_export[col], errors='coerce')
                    try:
                        raw_export[col] = raw_export[col].dt.tz_localize(None)
                    except:
                        pass

            raw_export.to_excel(writer, sheet_name="Raw Data", index=False)

            # --- HELPER FOR STATS ---
            def create_stats_df(df):
                try:
                    s_out = df["avg_output_rate"].std() if "avg_output_rate" in df else 0
                    s_time = df["avg_purge_mins"].std() if "avg_purge_mins" in df else 0
                    s_mat = df["avg_purge_mat"].std() if "avg_purge_mat" in df else 0
                    s_yld = df["avg_yield"].std() if "avg_yield" in df and not df["avg_yield"].isna().all() else 0
                except:
                    s_out = 0;
                    s_time = 0;
                    s_mat = 0;
                    s_yld = 0

                # Fill NaNs
                vals = [s_out, s_time, s_mat, s_yld]
                vals = [0 if pd.isna(x) else x for x in vals]

                return pd.DataFrame({
                    "Average Types": [
                        "Average Output/Hr", "Average Cleaning Time (min)",
                        "Average Kg Cleaning Material", "Average Yield %"
                    ],
                    "Std Deviation": vals
                })

            col_map = {
                "product_code": "Product Code", "machine_number": "Machine Number",
                "formula_number": "Formula No", "avg_output_rate": "Average Output/Hr",
                "avg_purge_mins": "Average Cleaning Time (min)",
                "avg_purge_mat": "Average Kg Cleaning Material",
                "avg_yield": "Average Yield %"
            }
            main_cols = list(col_map.values())

            # --- SHEET 2: VALID YIELDS ---
            if not summary_valid.empty:
                rep_ex = summary_valid.rename(columns=col_map)
                valid_cols = [c for c in main_cols if c in rep_ex.columns]
                rep_ex[valid_cols].to_excel(writer, sheet_name="Benchmark Report", index=False, startcol=0)

                stats_df = create_stats_df(summary_valid)
                stats_df.to_excel(writer, sheet_name="Benchmark Report", index=False, startcol=9)

            # --- SHEET 3: UNSPECIFIED YIELDS ---
            if not summary_invalid.empty:
                inv_ex = summary_invalid.rename(columns=col_map)
                valid_cols = [c for c in main_cols if c in inv_ex.columns]
                inv_ex[valid_cols].to_excel(writer, sheet_name="Unspecified Yields", index=False)

                stats_inv = create_stats_df(summary_invalid)
                stats_inv.to_excel(writer, sheet_name="Unspecified Yields", index=False, startcol=9)

            # --- FORMATTING ---
            wb = writer.book
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            center_align = Alignment(horizontal='center')

            target_sheets = []
            if "Benchmark Report" in wb.sheetnames: target_sheets.append(wb["Benchmark Report"])
            if "Unspecified Yields" in wb.sheetnames: target_sheets.append(wb["Unspecified Yields"])

            for ws in target_sheets:
                # Format Main Table
                for col_idx in range(1, 8):
                    ws.cell(row=1, column=col_idx).font = header_font
                    ws.cell(row=1, column=col_idx).fill = header_fill
                    ws.cell(row=1, column=col_idx).alignment = center_align

                # Format Stats Table
                for col_idx in range(10, 12):
                    ws.cell(row=1, column=col_idx).font = header_font
                    ws.cell(row=1, column=col_idx).fill = header_fill
                    ws.cell(row=1, column=col_idx).alignment = center_align

                # Number Formatting (0.00)
                for row in range(2, ws.max_row + 1):
                    for col in [4, 5, 6, 7]:  # D, E, F, G
                        cell = ws.cell(row=row, column=col)
                        if cell.value is not None: cell.number_format = '#,##0.00'

                # Stats rows
                for row in range(2, 6):
                    cell = ws.cell(row=row, column=11)  # K
                    if cell.value is not None: cell.number_format = '#,##0.00'

                # Widths
                from openpyxl.utils import get_column_letter
                for col_idx in range(1, 12):
                    col_letter = get_column_letter(col_idx)
                    ws.column_dimensions[col_letter].width = 22