# app/views/extruder_form/records/main.py

import decimal
import os
import traceback
from datetime import datetime, timedelta
from typing import Type, List

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.copier import WorksheetCopy

from PyQt6.QtCore import Qt, QTimer, pyqtSignal, QDate, QPoint, QThread
from PyQt6.QtGui import QAction, QColor, QBrush
from PyQt6.QtWidgets import (QWidget, QTableWidgetItem, QMessageBox, QApplication, QMenu, QDialog, QProgressDialog,
                             QFileDialog)
from sqlalchemy.orm import Session

from .bulk_export_dialog import BulkExportDialog
from .edit_dialog import ExtruderEditDialog
from .exporter import ExcelReportExporter
from .filter_dialog import FilterDialog
from .ops import ExtruderRecordsOperations
from .ui_setup import Ui_ExtruderRecordsList
from .view_dialog import ExtruderRecordViewDialog
from ..entry_form.widgets.error_dialog import ErrorDialog


class NumericTableWidgetItem(QTableWidgetItem):
    def __lt__(self, other):
        try:
            self_data = float(self.data(Qt.ItemDataRole.UserRole))
        except (ValueError, TypeError):
            self_data = 0.0
        try:
            other_data = float(other.data(Qt.ItemDataRole.UserRole))
        except (ValueError, TypeError):
            other_data = 0.0
        return self_data < other_data


class BulkExportWorker(QThread):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(str)
    error = pyqtSignal(str)

    def __init__(self, records: List, option: int, output_path: str, ops: ExtruderRecordsOperations, parent=None):
        super().__init__(parent)
        self.records = records
        self.option = option
        self.output_path = output_path
        self.ops = ops
        self._is_running = True

    def stop(self):
        self._is_running = False

    def run(self):
        try:
            exporter = ExcelReportExporter(self.ops)
            total_records = len(self.records)

            if self.option == BulkExportDialog.SEPARATE_FILES:
                self._export_to_separate_files(exporter, total_records)
            elif self.option == BulkExportDialog.SEPARATE_SHEETS:
                self._export_to_separate_sheets(exporter, total_records)
            elif self.option == BulkExportDialog.SINGLE_SHEET:
                self._export_to_single_sheet(exporter, total_records)

            if self._is_running:
                self.finished.emit("Bulk export completed successfully!")
        except Exception:
            self.error.emit(f"An unexpected error occurred during bulk export:\n\n{traceback.format_exc()}")

    def _export_to_separate_files(self, exporter, total):
        os.makedirs(self.output_path, exist_ok=True)
        for i, record in enumerate(self.records):
            if not self._is_running: return
            self.progress.emit(int((i / total) * 100), f"Exporting {record.ref_no} ({i + 1}/{total})...")
            filename = f"Report_{record.ref_no}_{record.lot_number}.xlsx".replace('/', '_').replace('\\', '_')
            file_path = os.path.join(self.output_path, filename)
            exporter.generate_single_report(record, file_path)
        self.progress.emit(100, "Finalizing...")

    def _export_to_separate_sheets(self, exporter, total):
        template_wb = openpyxl.load_workbook(ExcelReportExporter.TEMPLATE_PATH)
        template_sheet = template_wb.active

        output_wb = openpyxl.Workbook()
        output_wb.remove(output_wb.active)

        for i, record in enumerate(self.records):
            if not self._is_running: return
            self.progress.emit(int((i / total) * 100), f"Processing {record.ref_no} ({i + 1}/{total})...")

            sheet_title = f"Ref {record.ref_no}".replace('/', '_').replace('\\', '_')[:31]
            new_sheet = output_wb.create_sheet(title=sheet_title)

            for col_letter in [get_column_letter(c) for c in range(1, template_sheet.max_column + 1)]:
                new_sheet.column_dimensions[col_letter].width = template_sheet.column_dimensions[col_letter].width
            for row_idx, dim in template_sheet.row_dimensions.items():
                if dim.height is not None:
                    new_sheet.row_dimensions[row_idx].height = dim.height

            for row in template_sheet.iter_rows():
                for cell in row:
                    new_cell = new_sheet.cell(row=cell.row, column=cell.column)
                    new_cell.value = cell.value
                    if cell.has_style:
                        new_cell.font = cell.font.copy()
                        new_cell.border = cell.border.copy()
                        new_cell.fill = cell.fill.copy()
                        new_cell.number_format = cell.number_format
                        new_cell.protection = cell.protection.copy()
                        new_cell.alignment = cell.alignment.copy()

            for mc_range in template_sheet.merged_cells.ranges:
                new_sheet.merge_cells(str(mc_range))

            exporter.populate_sheet(new_sheet, record, overwrite_formulas=False)

        self.progress.emit(100, "Saving file...")
        output_wb.save(self.output_path)

    # def _export_to_single_sheet(self, exporter, total):
    #     output_wb = openpyxl.Workbook()
    #     output_sheet = output_wb.active
    #     output_sheet.title = "Bulk Report"
    #
    #     template_wb = openpyxl.load_workbook(ExcelReportExporter.TEMPLATE_PATH)
    #     template_sheet = template_wb.active
    #     for col_letter in [get_column_letter(i) for i in range(1, template_sheet.max_column + 1)]:
    #         output_sheet.column_dimensions[col_letter].width = template_sheet.column_dimensions[col_letter].width
    #
    #     template_row_heights = {i: dim.height for i, dim in template_sheet.row_dimensions.items()}
    #     default_row_height = template_sheet.sheet_format.defaultRowHeight
    #
    #     for i, record in enumerate(self.records):
    #         if not self._is_running: return
    #         self.progress.emit(int((i / total) * 100), f"Processing {record.ref_no} ({i + 1}/{total})...")
    #
    #         temp_wb = openpyxl.load_workbook(ExcelReportExporter.TEMPLATE_PATH)
    #         temp_sheet = temp_wb.active
    #
    #         exporter.populate_sheet(temp_sheet, record, row_offset=0, overwrite_formulas=True)
    #
    #         row_offset = i * ExcelReportExporter.REPORT_TOTAL_ROWS
    #         for row_idx in range(1, ExcelReportExporter.REPORT_TOTAL_ROWS + 2):
    #             height = template_row_heights.get(row_idx, default_row_height)
    #             if height is not None:
    #                 output_sheet.row_dimensions[row_idx + row_offset].height = height
    #
    #         for row in temp_sheet.iter_rows():
    #             for cell in row:
    #                 new_cell = output_sheet.cell(row=cell.row + row_offset, column=cell.column)
    #                 new_cell.value = cell.value
    #                 if cell.has_style:
    #                     new_cell.font = cell.font.copy()
    #                     new_cell.border = cell.border.copy()
    #                     new_cell.fill = cell.fill.copy()
    #                     new_cell.number_format = cell.number_format
    #                     new_cell.protection = cell.protection.copy()
    #                     new_cell.alignment = cell.alignment.copy()
    #
    #         for mc_range in temp_sheet.merged_cells.ranges:
    #             output_sheet.merge_cells(start_row=mc_range.min_row + row_offset, start_column=mc_range.min_col,
    #                                      end_row=mc_range.max_row + row_offset, end_column=mc_range.max_col)
    #
    #     self.progress.emit(100, "Saving file...")
    #     output_wb.save(self.output_path)

    def _export_to_single_sheet(self, exporter, total):
        output_wb = openpyxl.Workbook()
        output_sheet = output_wb.active
        output_sheet.title = "Bulk Report"

        # Setup columns based on template
        template_wb = openpyxl.load_workbook(ExcelReportExporter.TEMPLATE_PATH)
        template_sheet = template_wb.active
        for col_letter in [get_column_letter(i) for i in range(1, template_sheet.max_column + 1)]:
            output_sheet.column_dimensions[col_letter].width = template_sheet.column_dimensions[col_letter].width

        # We need to copy row heights carefully, but since heights vary per report,
        # we will do it dynamically inside the loop.

        current_write_row = 0

        for i, record in enumerate(self.records):
            if not self._is_running: return
            self.progress.emit(int((i / total) * 100), f"Processing {record.ref_no} ({i + 1}/{total})...")

            # 1. Load a fresh temp template for this record
            temp_wb = openpyxl.load_workbook(ExcelReportExporter.TEMPLATE_PATH)
            temp_sheet = temp_wb.active

            # 2. Populate it (offset 0). The exporter will insert rows into temp_sheet if needed.
            # extra_rows = exporter.populate_sheet(temp_sheet, record, row_offset=0, overwrite_formulas=True)
            # We don't strictly need the return value here because we just copy whatever the temp_sheet became.
            exporter.populate_sheet(temp_sheet, record, row_offset=0, overwrite_formulas=True)

            # 3. Copy the (potentially expanded) temp_sheet to the main output_sheet
            # We iterate over all rows in the temp sheet.
            max_row_in_temp = temp_sheet.max_row

            # Copy row dimensions
            for r in range(1, max_row_in_temp + 1):
                # Row height
                if r in temp_sheet.row_dimensions:
                    dim = temp_sheet.row_dimensions[r]
                    if dim.height is not None:
                        output_sheet.row_dimensions[current_write_row + r].height = dim.height

            # Copy cells
            for row in temp_sheet.iter_rows():
                for cell in row:
                    new_cell = output_sheet.cell(row=cell.row + current_write_row, column=cell.column)
                    new_cell.value = cell.value
                    if cell.has_style:
                        new_cell.font = cell.font.copy()
                        new_cell.border = cell.border.copy()
                        new_cell.fill = cell.fill.copy()
                        new_cell.number_format = cell.number_format
                        new_cell.protection = cell.protection.copy()
                        new_cell.alignment = cell.alignment.copy()

            # Copy Merged Cells
            for mc_range in temp_sheet.merged_cells.ranges:
                output_sheet.merge_cells(
                    start_row=mc_range.min_row + current_write_row,
                    start_column=mc_range.min_col,
                    end_row=mc_range.max_row + current_write_row,
                    end_column=mc_range.max_col
                )

            # 4. Update the write pointer for the next report
            # We add max_row_in_temp to stack them immediately after one another.
            # (Optionally add +1 if a gap is desired, but standard is usually continuous)
            current_write_row += max_row_in_temp

        self.progress.emit(100, "Saving file...")
        output_wb.save(self.output_path)

class ExtruderRecordsView(QWidget):
    data_changed = pyqtSignal()

    def __init__(self, session_factory: Type[Session], parent=None):
        super().__init__(parent)
        self.ui = Ui_ExtruderRecordsList()
        self.ui.setupUi(self)
        self.ops = ExtruderRecordsOperations(session_factory)

        self.bulk_export_worker = None

        self.filter_dialog = FilterDialog(session_factory, self)
        self.advanced_filters = {}
        self.is_loading = False
        self.is_deleted_color = QColor("#e0e0e0")

        self._setup_connections()
        self.load_initial_data()
        self._update_ui_for_view_mode()

        self.ui.table_widget.sortByColumn(1, Qt.SortOrder.DescendingOrder)

        css_path = os.path.join(os.path.dirname(__file__), "styles.css")
        if os.path.exists(css_path):
            with open(css_path, "r") as f: self.setStyleSheet(f.read())

    def _setup_connections(self):
        self.search_timer = QTimer(self)
        self.search_timer.setSingleShot(True)
        self.search_timer.timeout.connect(self.refresh_data)
        self.ui.search_input.textChanged.connect(lambda: self.search_timer.start(500))
        self.ui.date_from_input.dateChanged.connect(self.refresh_data)
        self.ui.date_to_input.dateChanged.connect(self.refresh_data)
        self.ui.advanced_filter_button.clicked.connect(self._open_filter_dialog)
        self.ui.clear_filters_button.clicked.connect(self._clear_all_filters)
        self.ui.restore_selected_button.clicked.connect(self._restore_selected_records)
        self.ui.show_only_deleted_checkbox.stateChanged.connect(self.refresh_data)
        self.ui.table_widget.doubleClicked.connect(self._view_record)
        self.ui.table_widget.customContextMenuRequested.connect(self._show_context_menu)

    def _open_filter_dialog(self):
        self.filter_dialog.populate_dropdowns()
        if self.filter_dialog.exec() == QDialog.DialogCode.Accepted:
            self.advanced_filters = self.filter_dialog.get_filters()
            self.refresh_data()

    def _clear_all_filters(self):
        self.ui.search_input.clear()
        self.ui.date_from_input.setDate(QDate.currentDate().addMonths(-1))
        self.ui.date_to_input.setDate(QDate.currentDate())
        self.advanced_filters.clear()
        self.refresh_data()

    def refresh_data(self):
        self._load_records()
        self._update_ui_for_view_mode()

    def load_initial_data(self):
        self._load_records()

    def _update_ui_for_view_mode(self):
        is_deleted_view = self.ui.show_only_deleted_checkbox.isChecked()
        self.ui.restore_selected_button.setVisible(is_deleted_view)

    def _restore_selected_records(self):
        selected_rows = set(item.row() for item in self.ui.table_widget.selectedItems())
        if not selected_rows:
            QMessageBox.information(self, "No Selection", "Please select one or more records to restore.")
            return
        ids_to_restore = [int(self.ui.table_widget.item(row, 0).text()) for row in selected_rows]
        count = len(ids_to_restore)
        reply = QMessageBox.question(self, 'Confirm Restore',
                                     f"Are you sure you want to restore {count} selected record(s)?",
                                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                                     QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            success = self.ops.restore_multiple_records(ids_to_restore)
            if success:
                QMessageBox.information(self, "Success", f"{count} record(s) have been restored.")
                self.refresh_data()
                self.data_changed.emit()
            else:
                QMessageBox.critical(self, "Error", "Failed to restore the selected records.")

    def _show_context_menu(self, position: QPoint):
        record_id, is_deleted = self._get_selected_record_info()
        if record_id is None: return
        context_menu = QMenu(self)
        if is_deleted:
            view_action = QAction("View Record", self)
            restore_action = QAction("Restore Record", self)
            view_action.triggered.connect(self._view_record)
            restore_action.triggered.connect(self._restore_record)
            context_menu.addAction(view_action)
            context_menu.addAction(restore_action)
        else:
            view_action = QAction("View Record", self)
            edit_action = QAction("Edit Record", self)
            delete_action = QAction("Delete Record", self)
            export_action = QAction("Export to Excel", self)
            view_action.triggered.connect(self._view_record)
            edit_action.triggered.connect(self._edit_record)
            delete_action.triggered.connect(self._delete_record)
            export_action.triggered.connect(self._export_record_to_excel)
            context_menu.addAction(view_action)
            context_menu.addAction(edit_action)
            context_menu.addSeparator()
            context_menu.addAction(export_action)
            record_count = self.ui.table_widget.rowCount()
            if record_count > 1:
                bulk_export_action = QAction(f"Bulk Export ({record_count} Records)...", self)
                bulk_export_action.triggered.connect(self._start_bulk_export)
                context_menu.addAction(bulk_export_action)
            context_menu.addSeparator()
            context_menu.addAction(delete_action)
        context_menu.exec(self.ui.table_widget.mapToGlobal(position))

    def _start_bulk_export(self):
        if self.bulk_export_worker:
            QMessageBox.warning(self, "Export in Progress", "A bulk export is already running.")
            return
        record_count = self.ui.table_widget.rowCount()
        dialog = BulkExportDialog(record_count, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            option = dialog.get_selected_option()

            QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
            try:
                record_ids = [int(self.ui.table_widget.item(row, 0).text()) for row in range(record_count)]
                records_to_sort = [self.ops.get_full_record_by_id(rid) for rid in record_ids]

                def get_min_start_time(record):
                    start_times = [out.datetime_start for out in record.extruder_outputs if out.datetime_start]
                    return min(start_times) if start_times else datetime.max

                sorted_records = sorted(
                    records_to_sort,
                    key=lambda r: (getattr(r.machine, 'name', ''), get_min_start_time(r))
                )
            finally:
                QApplication.restoreOverrideCursor()

            output_path = ""
            if option == BulkExportDialog.SEPARATE_FILES:
                output_path = QFileDialog.getExistingDirectory(self, "Select Folder to Save Reports")
            else:
                default_filename = f"Bulk_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
                output_path, _ = QFileDialog.getSaveFileName(self, "Save Bulk Report", default_filename,
                                                             "Excel Files (*.xlsx)")

            if not output_path:
                return

            self.progress = QProgressDialog("Starting bulk export...", "Cancel", 0, 100, self)
            self.progress.setWindowTitle("Bulk Exporting")
            self.progress.setWindowModality(Qt.WindowModality.WindowModal)

            self.bulk_export_worker = BulkExportWorker(sorted_records, option, output_path, self.ops)
            self.bulk_export_worker.progress.connect(self._update_bulk_progress)
            self.bulk_export_worker.finished.connect(self._on_bulk_export_finished)
            self.bulk_export_worker.error.connect(self._on_bulk_export_error)

            self.progress.canceled.connect(self.bulk_export_worker.stop)
            self.progress.show()

            self.bulk_export_worker.start()

    def _update_bulk_progress(self, value, text):
        self.progress.setValue(value)
        self.progress.setLabelText(text)

    def _on_bulk_export_finished(self, message):
        self.progress.setValue(100)
        QMessageBox.information(self, "Success", message)
        self.bulk_export_worker = None

    def _on_bulk_export_error(self, error_message):
        if hasattr(self, 'progress') and self.progress: self.progress.close()
        ErrorDialog("Bulk Export Error", "An error occurred during the bulk export.", details=error_message,
                    parent=self).exec()
        self.bulk_export_worker = None

    def _export_record_to_excel(self):
        if self.bulk_export_worker:
            QMessageBox.warning(self, "Export in Progress", "A bulk export is already running. Please wait.")
            return
        record_id, _ = self._get_selected_record_info()
        if record_id is None: return
        try:
            record = self.ops.get_full_record_by_id(record_id)
            if not record:
                QMessageBox.warning(self, "Not Found", "Could not retrieve the full record details for export.")
                return
        except Exception:
            ErrorDialog("Database Error", "Failed to fetch record data.", details=traceback.format_exc(),
                        parent=self).exec()
            return
        default_filename = f"Extruder_Report_{record.ref_no}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(self, "Save Excel Report", default_filename, "Excel Files (*.xlsx)")
        if not file_path:
            return
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            exporter = ExcelReportExporter(self.ops)
            exporter.generate_single_report(record, file_path)

            reply = QMessageBox.information(self, "Export Successful",
                                            f"Report successfully saved.\n\nDo you want to open the file?",
                                            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
            if reply == QMessageBox.StandardButton.Yes:
                os.startfile(file_path)
        except Exception:
            ErrorDialog("Export Error", "Failed to save the Excel file.", details=traceback.format_exc(),
                        parent=self).exec()
        finally:
            QApplication.restoreOverrideCursor()

    def _load_records(self):
        if self.is_loading: return
        self.is_loading = True
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        self.ui.table_widget.setSortingEnabled(False)
        self.ui.table_widget.setRowCount(0)
        try:
            all_filters = self.advanced_filters.copy()
            all_filters['search_term'] = self.ui.search_input.text()
            all_filters['date_from'] = self.ui.date_from_input.date().toPyDate()
            all_filters['date_to'] = self.ui.date_to_input.date().toPyDate()
            all_filters['show_only_deleted'] = self.ui.show_only_deleted_checkbox.isChecked()

            if search_term := self.ui.search_input.text():
                all_filters['ref_no_search'] = search_term

            records = self.ops.get_records_with_details(filters=all_filters)
            for record_object in records:
                self._add_record_to_table(record_object)
        except Exception:
            error_dialog = ErrorDialog("Database Error", "Could not load records.", details=traceback.format_exc(),
                                       parent=self)
            error_dialog.exec()
        finally:
            self.is_loading = False
            self.ui.table_widget.setSortingEnabled(True)
            QApplication.restoreOverrideCursor()

    def _add_record_to_table(self, record):
        row_pos = self.ui.table_widget.rowCount()
        self.ui.table_widget.insertRow(row_pos)

        total_output = sum(out.qty_output for out in record.extruder_outputs if out.qty_output)
        total_production_seconds = sum(
            (out.datetime_end - out.datetime_start).total_seconds() for out in record.extruder_outputs
            if out.datetime_start and out.datetime_end and out.datetime_end > out.datetime_start
        )
        total_production_hours = total_production_seconds / 3600.0
        output_per_hour = (total_output / decimal.Decimal(
            total_production_hours)) if total_production_hours > 0 else decimal.Decimal(0)
        total_purging_seconds = 0
        for p in record.purging_headers:
            if p.time_start and p.time_end:
                dummy_date = datetime.now().date()
                start_dt = datetime.combine(dummy_date, p.time_start)
                end_dt = datetime.combine(dummy_date, p.time_end)
                if end_dt < start_dt:
                    end_dt += timedelta(days=1)
                total_purging_seconds += (end_dt - start_dt).total_seconds()
        purging_hours, rem = divmod(total_purging_seconds, 3600)
        purging_minutes, _ = divmod(rem, 60)
        purging_time_str = f"{int(purging_hours):02}:{int(purging_minutes):02}"
        start_times = [out.datetime_start for out in record.extruder_outputs if out.datetime_start]
        end_times = [out.datetime_end for out in record.extruder_outputs if out.datetime_end]
        min_start_time = min(start_times) if start_times else None
        max_end_time = max(end_times) if end_times else None

        cell_data = {
            0: (str(record.id), record.id),
            1: (record.created_at.strftime("%Y-%m-%d %H:%M") if record.created_at else "N/A",
                record.created_at.timestamp() if record.created_at else 0),
            2: (str(record.ref_no or ''), record.ref_no),
            3: (getattr(record.machine, 'name', 'N/A'), None),
            4: (record.product_code, None),
            5: (record.lot_number, None),
            6: (min_start_time.strftime("%Y-%m-%d %H:%M") if min_start_time else "N/A",
                min_start_time.timestamp() if min_start_time else 0),
            7: (max_end_time.strftime("%Y-%m-%d %H:%M") if max_end_time else "N/A",
                max_end_time.timestamp() if max_end_time else 0),
            8: (f"{output_per_hour:.2f}", float(output_per_hour)),
            9: (f"{record.target_output_per_hour or 0:.2f}", float(record.target_output_per_hour or 0)),
            10: (f"{total_output or 0:.2f}", float(total_output or 0)),
            11: (", ".join([p.product_code for p in record.purging_headers if p.product_code]) or "N/A", None),
            12: (purging_time_str, total_purging_seconds),
            13: (", ".join([f"{p.employee.first_name} {p.employee.last_name}" for p in record.extruder_personnels if
                            p.employee]) or "N/A", None),
        }

        is_deleted_view = self.ui.show_only_deleted_checkbox.isChecked()
        for col, (text, numeric_val) in cell_data.items():
            item = NumericTableWidgetItem(text) if numeric_val is not None else QTableWidgetItem(text)
            if numeric_val is not None:
                item.setData(Qt.ItemDataRole.UserRole, numeric_val)
            if is_deleted_view:
                item.setBackground(QBrush(QColor("#e0e0e0")))
            self.ui.table_widget.setItem(row_pos, col, item)

    def _get_selected_record_info(self):
        selected_rows = self.ui.table_widget.selectionModel().selectedRows()
        if not selected_rows: return None, False
        row = selected_rows[0].row()
        record_id = int(self.ui.table_widget.item(row, 0).text())
        is_deleted = self.ui.show_only_deleted_checkbox.isChecked()
        return record_id, is_deleted

    def _delete_record(self):
        record_id, is_deleted = self._get_selected_record_info()
        if record_id is None or is_deleted: return
        reply = QMessageBox.question(self, 'Confirm Deletion', f"Delete record ID: {record_id}?",
                                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                                     QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            if self.ops.soft_delete_record(record_id):
                self.refresh_data()
                self.data_changed.emit()
            else:
                QMessageBox.critical(self, "Error", "Failed to delete the record.")

    def _restore_record(self):
        record_id, is_deleted = self._get_selected_record_info()
        if record_id is None or not is_deleted: return
        reply = QMessageBox.question(self, 'Confirm Restore',
                                     f"Are you sure you want to restore record ID: {record_id}?",
                                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                                     QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            if self.ops.restore_record(record_id):
                QMessageBox.information(self, "Success", f"Record {record_id} has been restored.")
                self.refresh_data()
                self.data_changed.emit()
            else:
                QMessageBox.critical(self, "Error", "Failed to restore the record.")

    def _view_record(self):
        record_id, is_deleted = self._get_selected_record_info()
        if record_id is None: return
        full_data = self.ops.get_full_record_by_id(record_id)
        if full_data:
            view_dialog = ExtruderRecordViewDialog(self)
            view_dialog.populate_data(full_data)
            view_dialog.exec()
        else:
            QMessageBox.warning(self, "Not Found", "Could not retrieve full record details.")

    def _edit_record(self):
        record_id, is_deleted = self._get_selected_record_info()
        if record_id is None or is_deleted: return
        full_data = self.ops.get_full_record_by_id(record_id)
        if full_data:
            edit_dialog = ExtruderEditDialog(self.ops.Session, full_data, self)
            if edit_dialog.exec() == QDialog.DialogCode.Accepted:
                self.refresh_data()
                self.data_changed.emit()
        else:
            QMessageBox.warning(self, "Not Found", "Could not retrieve full record details for editing.")