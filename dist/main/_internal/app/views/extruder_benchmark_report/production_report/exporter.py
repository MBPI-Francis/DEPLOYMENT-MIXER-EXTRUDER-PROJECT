import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment


class DynamicProductionExporter:
    def export(self, df: pd.DataFrame, visible_columns: list, filepath: str):
        """
        visible_columns: List of tuples [('Internal_Col_Name', 'Display Name'), ...]
        """
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:

            # 1. Construct Export DataFrame
            export_df = pd.DataFrame()

            # List of columns known to potentially contain datetimes
            # We will check these specifically to be safe
            potential_date_cols = [
                "created_at", "time_start", "time_end",
                "purging_start_time", "purging_end_time"
            ]

            for internal_col, display_name in visible_columns:
                if internal_col in df.columns:
                    col_data = df[internal_col]

                    # Logic: Check if column is in our known date list OR if pandas thinks it's a date
                    is_date_col = (internal_col in potential_date_cols) or \
                                  pd.api.types.is_datetime64_any_dtype(col_data)

                    if is_date_col:
                        try:
                            # 1. Force convert to pandas datetime (handles mixed types/objects)
                            # errors='coerce' turns bad data into NaT (empty)
                            series = pd.to_datetime(col_data, errors='coerce')

                            # 2. Remove Timezone info (make it "naive")
                            # This handles both tz-aware and already-naive data gracefully
                            if series.dt.tz is not None:
                                export_df[display_name] = series.dt.tz_localize(None)
                            else:
                                export_df[display_name] = series

                        except Exception:
                            # If conversion fails, fallback to string representation
                            export_df[display_name] = col_data.astype(str)
                    else:
                        # Just copy data as is
                        export_df[display_name] = col_data
                else:
                    export_df[display_name] = ""

                    # 2. Write to Sheet
            sheet_name = "Production Report"
            export_df.to_excel(writer, sheet_name=sheet_name, index=False)

            # 3. Formatting
            wb = writer.book
            ws = wb[sheet_name]

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            center_align = Alignment(horizontal='center')

            # Format Headers
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

            # Auto-width
            from openpyxl.utils import get_column_letter
            for col_idx, column_cells in enumerate(ws.columns, 1):
                try:
                    # Calculate width based on max string length
                    max_len = 0
                    for cell in column_cells:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))

                    # Cap the width to prevent massive columns
                    final_width = min(max(max_len + 5, 15), 50)
                    ws.column_dimensions[get_column_letter(col_idx)].width = final_width
                except:
                    ws.column_dimensions[get_column_letter(col_idx)].width = 20