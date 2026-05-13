# app/views/extruder_form/records/exporter.py

import os
import openpyxl
from copy import copy
from datetime import datetime, timedelta
from decimal import Decimal
import pandas as pd
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.cell import get_column_letter

from .ops import ExtruderRecordsOperations


class ExcelReportExporter:
    """
    Handles populating the extruder report Excel template.
    Strictly follows the algorithm:
    1. Insert Data Rows.
    2. Format New Data Rows (Styles).
    3. FORCE REPAIR Footer Merges (Last 7 Rows).
    4. SEVER Footer/Data Boundaries (Prevent overlaps).
    5. FINAL PASS: Force Data Merges (I:J) on all data rows.
    """
    TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "extruder_report_format.xlsx")

    # Template constants
    DEFAULT_DATA_ROWS = 13
    DATA_START_ROW = 12
    FOOTER_START_ROW = 25  # The row where the footer initially starts (Row 25)
    TEMPLATE_LAST_DATA_ROW = 24  # The row index of the last data row (Row 24)

    def __init__(self, ops_controller: ExtruderRecordsOperations):
        self.controller = ops_controller

    def generate_single_report(self, record_object, output_path: str):
        """Creates and saves a report for a single record to a new file."""
        workbook = openpyxl.load_workbook(self.TEMPLATE_PATH)
        sheet = workbook.active
        self.populate_sheet(sheet, record_object, overwrite_formulas=False)
        workbook.save(output_path)

    def populate_sheet(self, sheet: Worksheet, record_object, row_offset: int = 0, overwrite_formulas: bool = False):
        """
        Fills a given worksheet with data.
        """
        # 1. Fetch and Prepare Data
        sorted_outputs = sorted(record_object.extruder_outputs,
                                key=lambda x: (x.datetime_start is None, x.datetime_start))

        valid_materials = []
        if record_object.production_id:
            prod_ids = [pid.strip() for pid in record_object.production_id.split(';') if pid.strip()]
            if prod_ids:
                raw_materials = self.controller.get_aggregated_materials_for_production_ids(prod_ids)
                valid_materials = [
                    (mat, qty) for mat, qty in raw_materials
                    if mat and str(mat).strip() and (qty or 0) > 0
                ]

        lot_numbers = []
        if record_object.lot_number:
            lot_numbers = [lot.strip() for lot in record_object.lot_number.split(';') if lot.strip()]

        # 2. Calculate if we need to insert rows
        n_outputs = len(sorted_outputs)
        n_materials = len(valid_materials)
        n_lots = len(lot_numbers)

        max_data_rows = max(n_outputs, n_materials, n_lots)
        extra_rows = max(0, max_data_rows - self.DEFAULT_DATA_ROWS)

        # 3. Dynamic Row Insertion
        insert_pos = self.FOOTER_START_ROW + row_offset
        source_row_idx = self.TEMPLATE_LAST_DATA_ROW + row_offset

        if extra_rows > 0:
            # --- A. Prepare Source Row (Row 24) ---
            saved_bottom_border_style = None

            for col in [9, 10]:  # I and J
                cell_24 = sheet.cell(row=source_row_idx, column=col)
                if cell_24.has_style:
                    if col == 9:
                        saved_bottom_border_style = copy(cell_24.border.bottom)

                    old_bd = cell_24.border
                    new_bd = Border(
                        left=old_bd.left, right=old_bd.right, top=old_bd.top,
                        bottom=Side(style=None),
                        diagonal=old_bd.diagonal, outline=old_bd.outline,
                        vertical=old_bd.vertical, horizontal=old_bd.horizontal
                    )
                    cell_24.border = new_bd

            # --- B. Insert Rows ---
            sheet.insert_rows(insert_pos, amount=extra_rows)

            # --- C. Clean Ghost Merges ---
            new_rows_start = insert_pos
            new_rows_end = insert_pos + extra_rows - 1
            merges_to_remove = []
            for merge_range in sheet.merged_cells.ranges:
                if (merge_range.min_row >= new_rows_start and
                        merge_range.max_row <= new_rows_end):
                    merges_to_remove.append(merge_range)
            for m in merges_to_remove: sheet.merged_cells.remove(m)

            # --- D. Format New Rows (Styles Only) ---
            for i in range(extra_rows):
                target_row_idx = insert_pos + i

                # Copy Height
                if source_row_idx in sheet.row_dimensions:
                    sheet.row_dimensions[target_row_idx].height = sheet.row_dimensions[source_row_idx].height

                # Copy Cell Styles
                for col in range(1, sheet.max_column + 1):
                    source_cell = sheet.cell(row=source_row_idx, column=col)
                    target_cell = sheet.cell(row=target_row_idx, column=col)

                    if source_cell.has_style:
                        target_cell.font = copy(source_cell.font)
                        target_cell.fill = copy(source_cell.fill)
                        target_cell.number_format = source_cell.number_format
                        target_cell.alignment = copy(source_cell.alignment)
                        target_cell.protection = copy(source_cell.protection)

                        # Border Handling
                        if col in [9, 10]:
                            src_bd = source_cell.border
                            new_border = Border(
                                left=src_bd.left, right=src_bd.right,
                                top=Side(style=None),
                                bottom=Side(style=None),
                                diagonal=src_bd.diagonal, outline=src_bd.outline,
                                vertical=src_bd.vertical, horizontal=src_bd.horizontal
                            )
                            target_cell.border = new_border
                        else:
                            target_cell.border = copy(source_cell.border)

            # --- E. Restore Bottom Border to the Last Data Row ---
            last_data_row = new_rows_end
            if saved_bottom_border_style:
                for col in [9, 10]:
                    last_cell = sheet.cell(row=last_data_row, column=col)
                    old_bd = last_cell.border
                    new_bd_last = Border(
                        left=old_bd.left, right=old_bd.right, top=old_bd.top,
                        bottom=saved_bottom_border_style,
                        diagonal=old_bd.diagonal, outline=old_bd.outline,
                        vertical=old_bd.vertical, horizontal=old_bd.horizontal
                    )
                    last_cell.border = new_bd_last

        # 4. Define Indices
        current_data_end_row = self.TEMPLATE_LAST_DATA_ROW + row_offset + extra_rows
        footer_offset = row_offset + extra_rows

        # 5. REPAIR FOOTER MERGES (7 Rows)
        self._repair_footer_merges(sheet, self.FOOTER_START_ROW + footer_offset)

        # 6. SEVER DATA-FOOTER BOUNDARY
        # Ensure no vertical merges cross from Data (last_row) into Footer (first_row)
        self._sever_vertical_merges(sheet, current_data_end_row)

        # 7. FINAL PASS: FORCE DATA MERGES
        # We iterate through ALL data rows (start to end) and enforce the I:J merge.
        # This fixes the "last 3 rows unmerged" issue.
        data_start_abs = self.DATA_START_ROW + row_offset
        for r in range(data_start_abs, current_data_end_row + 1):
            # A:C
            self._safe_merge(sheet, r, 1, r, 3)
            # D:E
            self._safe_merge(sheet, r, 4, r, 5)
            # I:J
            self._safe_merge(sheet, r, 9, r, 10)

        # 8. Populate Data
        self._fill_header_info(sheet, record_object, row_offset)
        self._fill_summary_and_calculations(sheet, record_object, row_offset)
        self._fill_output_log(sheet, sorted_outputs, row_offset, footer_offset, overwrite_formulas,
                              current_data_end_row)
        self._fill_materials_and_lots(sheet, valid_materials, lot_numbers, row_offset, footer_offset,
                                      overwrite_formulas, current_data_end_row)
        self._fill_purging_info(sheet, record_object, footer_offset)
        self._fill_footer_info(sheet, record_object, footer_offset)

    def _safe_merge(self, sheet, r1, c1, r2, c2):
        """Helper to merge without crashing if already merged."""
        try:
            sheet.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        except:
            pass  # Already merged or conflict, usually fine as we cleaned before

    def _sever_vertical_merges(self, sheet, boundary_row):
        """
        Finds any merge that spans across `boundary_row` and `boundary_row + 1`.
        Removes it to prevent Data rows merging into Footer rows.
        """
        merges_to_remove = []
        for merge in sheet.merged_cells.ranges:
            if merge.min_row <= boundary_row and merge.max_row > boundary_row:
                merges_to_remove.append(merge)

        for m in merges_to_remove:
            sheet.merged_cells.remove(m)

    def _repair_footer_merges(self, sheet: Worksheet, footer_start_row: int):
        """
        Explicitly applies the required merges for the 7 footer rows.
        """

        def force_merge_rect(min_row, min_col, max_row, max_col):
            # 1. Identify overlaps
            overlaps = []
            for merge in sheet.merged_cells.ranges:
                if not (merge.max_col < min_col or merge.min_col > max_col or
                        merge.max_row < min_row or merge.min_row > max_row):
                    overlaps.append(merge)
            # 2. Remove overlaps
            for merge in overlaps:
                sheet.merged_cells.remove(merge)
            # 3. Apply new merge
            sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)

        # Relative Rows (0 to 6)
        r1 = footer_start_row  # Total Output
        r2 = footer_start_row + 1  # Remarks Start
        r3 = footer_start_row + 2
        r4 = footer_start_row + 3  # Resin Details | Operators (Start)
        r5 = footer_start_row + 4  # Resin Used    | Operators (End)
        r6 = footer_start_row + 5  # Siever        | Supervisor
        r7 = footer_start_row + 6  # Palletizer

        # --- ROW 1 (Totals) ---
        force_merge_rect(r1, 1, r1, 5)  # A:E

        # Clean I:J overlap on Row 1
        overlaps_ij = []
        for merge in sheet.merged_cells.ranges:
            if not (merge.max_col < 9 or merge.min_col > 10):
                # Check vertical intersection with Row 1
                if merge.min_row <= r1 <= merge.max_row:
                    overlaps_ij.append(merge)
        for m in overlaps_ij: sheet.merged_cells.remove(m)

        # Clean M:N overlap on Row 1 (Fixing the merged Total Input issue)
        overlaps_mn = []
        for merge in sheet.merged_cells.ranges:
            if not (merge.max_col < 13 or merge.min_col > 14):
                if merge.min_row <= r1 <= merge.max_row:
                    overlaps_mn.append(merge)
        for m in overlaps_mn: sheet.merged_cells.remove(m)

        # --- REMARKS BLOCK (Col G-J, Row 2 to 7) ---
        force_merge_rect(r2, 7, r7, 10)

        # --- LEFT SIDE MERGES (B:D or B:E) ---
        force_merge_rect(r2, 2, r2, 4)  # Row 2: B:D
        force_merge_rect(r3, 2, r3, 4)  # Row 3: B:D
        force_merge_rect(r4, 2, r4, 5)  # Row 4: B:E
        force_merge_rect(r5, 2, r5, 4)  # Row 5: B:D
        force_merge_rect(r6, 2, r6, 4)  # Row 6: B:D
        force_merge_rect(r7, 2, r7, 4)  # Row 7: B:D

        # --- RIGHT SIDE MERGES (Personnel) ---
        force_merge_rect(r4, 12, r5, 12)  # L28:L29 (Vertical)
        force_merge_rect(r4, 13, r5, 14)  # M28:N29 (Block)
        force_merge_rect(r6, 13, r6, 14)  # M30:N30
        force_merge_rect(r7, 13, r7, 14)

    def _fill_header_info(self, sheet, record, row_offset):
        machine_name = getattr(record.machine, 'name', 'N/A')
        sheet[f'F{5 + row_offset}'].value = f"Extruder Machine No. {machine_name}"

    def _fill_summary_and_calculations(self, sheet, record, row_offset):
        sheet[f'A{8 + row_offset}'].value = getattr(record.machine, 'name', 'N/A')
        sheet[f'B{8 + row_offset}'].value = record.qty_order
        sheet[f'C{8 + row_offset}'].value = record.customer
        sheet[f'F{8 + row_offset}'].value = record.product_code
        sheet[f'G{9 + row_offset}'].value = record.qty_produced

        total_output = sum(out.qty_output for out in record.extruder_outputs if out.qty_output)
        total_seconds = sum(
            (out.datetime_end - out.datetime_start).total_seconds()
            for out in record.extruder_outputs
            if out.datetime_start and out.datetime_end and out.datetime_end > out.datetime_start
        )
        total_hours_decimal = Decimal(total_seconds) / Decimal(3600)

        sheet[f'H{9 + row_offset}'].value = float(total_hours_decimal.quantize(Decimal('0.01')))
        output_per_hour = (Decimal(total_output) / total_hours_decimal) if total_hours_decimal > 0 else Decimal(0)
        sheet[f'I{9 + row_offset}'].value = float(output_per_hour.quantize(Decimal('0.01')))
        sheet[f'J{9 + row_offset}'].value = float(record.target_output_per_hour or 0)
        sheet[f'K{9 + row_offset}'].value = total_output

        qty_produced_dec = Decimal(record.qty_produced or 0)
        loss = max(Decimal(0), qty_produced_dec - Decimal(total_output))
        denominator = Decimal(total_output) + loss
        output_percentage = (Decimal(total_output) / denominator * 100) if denominator > 0 else Decimal(0)
        loss_percentage = (loss / denominator * 100) if denominator > 0 else Decimal(0)

        sheet[f'L{9 + row_offset}'].value = f"{output_percentage:.2f}%"
        sheet[f'M{9 + row_offset}'].value = loss
        sheet[f'N{9 + row_offset}'].value = f"{loss_percentage:.2f}%"

    def _fill_output_log(self, sheet, sorted_outputs, row_offset, footer_offset, overwrite_formulas,
                         current_data_end_row):
        total_duration = timedelta()
        for i, output in enumerate(sorted_outputs):
            row = self.DATA_START_ROW + i + row_offset
            if output.datetime_start:
                sheet[f'A{row}'].value = output.datetime_start.strftime("%d-%b-%Y %H:%M")
            if output.datetime_end:
                sheet[f'D{row}'].value = output.datetime_end.strftime("%d-%b-%Y %H:%M")
            if output.datetime_start and output.datetime_end:
                duration = output.datetime_end - output.datetime_start
                total_duration += duration
                hours, rem = divmod(duration.total_seconds(), 3600)
                minutes, _ = divmod(rem, 60)
                sheet[f'F{row}'].value = f"{int(hours):02d}:{int(minutes):02d}"
            output_cell = sheet[f'G{row}']
            output_cell.value = float(output.qty_output or 0)
            output_cell.number_format = '0.00'

        total_output_cell = sheet[f'G{25 + footer_offset}']
        if overwrite_formulas:
            total_output = sum(float(o.qty_output or 0) for o in sorted_outputs)
            total_output_cell.value = total_output
        else:
            start_row = self.DATA_START_ROW + row_offset
            total_output_cell.value = f"=SUM(G{start_row}:G{current_data_end_row})"
        total_output_cell.number_format = '#,##0.00'
        total_seconds = int(total_duration.total_seconds())
        hours, rem = divmod(total_seconds, 3600)
        minutes, seconds = divmod(rem, 60)
        sheet[f'F{25 + footer_offset}'].value = f"{hours:02d}:{minutes:02d}:{seconds:02d}"

    def _fill_materials_and_lots(self, sheet, valid_materials, lot_numbers, row_offset, footer_offset,
                                 overwrite_formulas, current_data_end_row):
        start_row = self.DATA_START_ROW + row_offset
        total_lot_qty_prod = 0.0

        # Fill Materials
        for i, (mat_code, total_qty) in enumerate(valid_materials):
            current_row = start_row + i
            mat_cell = sheet[f'I{current_row}']
            mat_cell.value = mat_code
            mat_cell.alignment = Alignment(horizontal='left', vertical='center')

            qty_cell = sheet[f'K{current_row}']
            qty_cell.value = float(total_qty or 0)
            qty_cell.alignment = Alignment(horizontal='right', vertical='center')
            qty_cell.number_format = '#,##0.0000000'

        lot_quantities_map = {}
        if lot_numbers:
            lot_quantities_map = self.controller.get_produced_quantities_for_lots(lot_numbers)

        for i, lot in enumerate(lot_numbers):
            row = start_row + i
            sheet[f'M{row}'].value = lot
            qty_prod = lot_quantities_map.get(lot, Decimal('0.00'))
            qty_prod_float = float(qty_prod or 0)
            total_lot_qty_prod += qty_prod_float
            qty_cell = sheet[f'N{row}']
            qty_cell.value = qty_prod_float
            qty_cell.alignment = Alignment(horizontal='right', vertical='center')
            qty_cell.number_format = '#,##0.00'

        total_input_cell = sheet[f'K{25 + footer_offset}']
        total_lot_qty_cell = sheet[f'N{25 + footer_offset}']

        data_start_idx = self.DATA_START_ROW + row_offset

        if overwrite_formulas:
            total_input_qty = sum(float(total_qty or 0) for _, total_qty in valid_materials)
            total_input_cell.value = total_input_qty
            total_lot_qty_cell.value = total_lot_qty_prod
        else:
            total_input_cell.value = f"=SUM(K{data_start_idx}:K{current_data_end_row})"
            total_lot_qty_cell.value = f"=SUM(N{data_start_idx}:N{current_data_end_row})"

        total_input_cell.number_format = '#,##0.00'
        total_lot_qty_cell.number_format = '#,##0.00'
        total_lot_qty_cell.alignment = Alignment(horizontal='right', vertical='center')

    def _fill_purging_info(self, sheet, record, footer_offset):
        if not record.purging_headers: return
        header = record.purging_headers[0]

        # Row 2 (B:D)
        sheet[f'B{26 + footer_offset}'].value = header.product_code

        # Row 3 (B:D)
        start_str = header.time_start.strftime("%H:%M") if header.time_start else ""
        end_str = header.time_end.strftime("%H:%M") if header.time_end else ""
        sheet[f'B{27 + footer_offset}'].value = f"{start_str} – {end_str}"

        if header.time_start and header.time_end:
            dummy_date = datetime.now().date()
            start_dt = datetime.combine(dummy_date, header.time_start)
            end_dt = datetime.combine(dummy_date, header.time_end)
            if end_dt < start_dt: end_dt += timedelta(days=1)
            duration = end_dt - start_dt
            total_seconds = int(duration.total_seconds())
            hours, rem = divmod(total_seconds, 3600)
            minutes, _ = divmod(rem, 60)
            sheet[f'E{27 + footer_offset}'].value = f"{int(hours):02d}:{int(minutes):02d}"

        # Row 4 (B:E) - Resin Details
        resin_details_str = ", ".join(
            [f"{getattr(d.resin, 'abbreviation', 'N/A')} = {d.qty}" for d in header.purging_details])
        sheet[f'B{28 + footer_offset}'].value = resin_details_str

        # Row 5 (B:D) - Resin Used Abbreviation
        resin_abbrev = getattr(header.resin_used, 'abbreviation', 'N/A')
        sheet[f'B{29 + footer_offset}'].value = resin_abbrev

        # Row 6 (B:D) - Siever
        sheet[f'B{30 + footer_offset}'].value = header.siever_used

        # Row 7 (B:D) - Palletizer
        sheet[f'B{31 + footer_offset}'].value = header.palletizer_used

    def _fill_footer_info(self, sheet, record, footer_offset):
        remarks_text = record.remarks or ""
        if record.machine_details and record.machine_details.is_vacuum_on:
            if remarks_text:
                remarks_text = f"VACUUM ON\n{remarks_text}"
            else:
                remarks_text = "VACUUM ON"
        sheet[f'G{26 + footer_offset}'].value = remarks_text
        sheet[f'G{26 + footer_offset}'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        operators, supervisors = [], []
        for p in record.extruder_personnels:
            if not p.employee or not p.position: continue
            full_name = f"{p.employee.first_name} {p.employee.last_name}"
            pos_name = p.position.name.lower()
            if "supervisor" in pos_name:
                supervisors.append(full_name)
            elif "operator" in pos_name:
                operators.append(full_name)
            elif "reliever" in pos_name:
                operators.append(f"{full_name} (Reliever)")

        # Operators -> Row 4 (Merged M28:N29)
        sheet[f'M{28 + footer_offset}'].value = ", ".join(operators) if operators else "N/A"
        sheet[f'M{28 + footer_offset}'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # Supervisors -> Row 6 (Merged M30:N30)
        sheet[f'M{30 + footer_offset}'].value = ", ".join(supervisors) if supervisors else "N/A"



class ExtruderListExporter:
    """
    Exports the current list view of Extruder Records to Excel,
    including the dynamic summary box at the bottom.
    """

    # Updated Columns List with new fields
    COLUMNS = [
        "Date Encoded",
        "Date Completed",
        "Reference No",
        "Machine Name",
        "Product Code",
        "Formula No",
        "Lot Number",
        "Customer",
        "Extrusion Time Start",
        "Extrusion Time End",
        "Extrusion Duration (hh:mm)",
        "Target Output/hr (kg)",
        "Actual Output/hr (kg)",
        "Total Output (kg)",
        "Purging To Code",
        "Cleaning Material",
        "Total Cleaning QTY (kg)",
        "Purging Duration (hh:mm)",
        "Operators"
    ]

    def export_list(self, data_list: list, summary_stats: dict, filepath: str):
        # 1. Create DataFrame
        df = pd.DataFrame(data_list)

        # Ensure column order & existence
        for col in self.COLUMNS:
            if col not in df.columns:
                df[col] = ""

        df = df[self.COLUMNS]

        # 2. Write to Excel
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            sheet_name = "Extruder Records"
            df.to_excel(writer, sheet_name=sheet_name, index=False)

            # 3. Access Workbook/Worksheet for formatting
            workbook = writer.book
            worksheet = writer.sheets[sheet_name]

            # --- STYLES DEFINITION ---
            header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
            header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")  # Dark Blue

            data_font = Font(name="Arial", size=10)

            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )

            # --- FORMAT MAIN DATA TABLE ---

            # A. Format Headers (Row 1)
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border

            # B. Format Data Rows
            # Define which columns are numeric for formatting
            numeric_cols_indices = []
            for idx, col_name in enumerate(self.COLUMNS):
                if any(x in col_name for x in ['Output', 'QTY', 'per/hr', 'Loss', 'Decimal']):
                    numeric_cols_indices.append(idx + 1)  # 1-based index

            for row in worksheet.iter_rows(min_row=2, max_row=len(df) + 1):
                for cell in row:
                    cell.font = data_font
                    cell.border = thin_border

                    # Apply alignment and number format
                    if cell.column in numeric_cols_indices:
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        # Check if value is actually a number before formatting
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0.00'
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

            # --- AUTO-FIT COLUMNS ---
            for i, column in enumerate(worksheet.columns):
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if cell.value:
                            val_len = len(str(cell.value))
                            if val_len > max_length:
                                max_length = val_len
                    except:
                        pass

                # Add some padding
                adjusted_width = min(max_length + 4, 60)  # Cap width at 60
                worksheet.column_dimensions[column_letter].width = adjusted_width

            # --- APPEND SUMMARY BOX ---
            last_row = len(df) + 1
            start_summary_row = last_row + 3

            # Helper to write summary row
            def write_summary_row(row_offset, label, value_key):
                r = start_summary_row + row_offset

                # Label Cell (Col A / 1)
                lbl_cell = worksheet.cell(row=r, column=1, value=label)
                lbl_cell.font = header_font
                lbl_cell.fill = header_fill
                lbl_cell.border = thin_border
                lbl_cell.alignment = Alignment(horizontal='left')

                # Value Cell (Col B / 2)
                val_cell = worksheet.cell(row=r, column=2, value=summary_stats.get(value_key, "0.00"))
                val_cell.font = data_font
                val_cell.border = thin_border
                val_cell.alignment = Alignment(horizontal='right')

                # Try to apply number format if it's not a time string (has ':')
                val_str = str(val_cell.value)
                if ':' not in val_str:
                    try:
                        # Attempt to convert to float to ensure number format works
                        val_cell.value = float(str(val_cell.value).replace(',', ''))
                        val_cell.number_format = '#,##0.00'
                    except:
                        pass  # Keep as string/formatted string

            # Write the 6 Rows
            write_summary_row(0, "Total Output Qty:", 'total_output')
            write_summary_row(1, "Total Processing Duration:", 'proc_duration_str')
            write_summary_row(2, "Processing (Excel Decimal):", 'proc_decimal')
            write_summary_row(3, "Total Purging Qty:", 'total_waste')
            write_summary_row(4, "Total Purging Duration:", 'waste_duration_str')
            write_summary_row(5, "Purging (Excel Decimal):", 'waste_decimal')

            # Ensure Column A and B are wide enough for the summary
            current_width_a = worksheet.column_dimensions['A'].width
            current_width_b = worksheet.column_dimensions['B'].width

            if current_width_a < 30: worksheet.column_dimensions['A'].width = 30
            if current_width_b < 20: worksheet.column_dimensions['B'].width = 20